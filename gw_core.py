"""
gw_core.py — GW Material Intake Form engine (Spare Parts / Z002).

Pure deterministic logic. No AI, no network, no disk writes: every function
takes and returns bytes or plain Python objects.

Rules here are the ones confirmed in the gw-intake-form skill as of
2026-09-28. If this file and that skill ever disagree, the skill wins and
this file gets fixed.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from datetime import date

import openpyxl
from openpyxl.styles import PatternFill

NO_FILL = PatternFill(fill_type=None)

INCH_TO_CM = 2.54
LB_TO_KG = 0.453592

CONFIG_VERSION = "2026-09-28"

# ---------------------------------------------------------------------------
# Plant configuration
# ---------------------------------------------------------------------------

EU_TAX_BY_COL = {56: "NL", 57: "MWST", 58: 1, 59: "IT", 60: "MWST", 61: 1, 107: "X"}
GB_TAX_BY_COL = {56: "GB", 57: "MWST", 58: 1, 107: "X"}

PLANT_CONFIG = {
    "US30": {
        "tab": "US",
        "length_unit": '"',
        "weight_unit": "LB",
        "fixed_by_col": {107: "X"},
        "fixed": {
            "MARA-MBRSH": "M", "MARA-MTART": "Z002", "MARC-TWERK": "US30",
            "MVKE-VKORG": "US30", "MVKE-VTWEG": "00", "MARA-SPART": "04",
            "MARA-MEINS": "EA", "MARA-MSTAE": "03", "MARM-MEINH": "EA",
            "MARM-UMREN": 1, "MARM-UMREZ": 1, "MARM-VOLEH": "FT3",
            "MVKE-DWERK": "US30", "MLAN-ALAND": "US", "TATYP": "ZTXJ",
            "MLAN-TAXM1": "1", "MVKE-KTGRM": "20", "MARA-MTPOS_MARA": "NORM",
            "MVKE-MTPOS": "NORM", "MARA-TRAGR": "0001", "MARC-LADGR": "0001",
            "MARA-VABME": 1, "MARC-EKGRP": "P07", "MARC-WEBAZ": 4,
            "MARC-MMSTA": "03", "MARC-DISMM": "PD", "MARC-DISPO": "Z01",
            "MARC-DISLS": "MB", "MARC-BESKZ": "F", "MARC-LGPRO": "U302",
            "MARC-LGFSB": "U302",
            "MARC -PLIFZ": 87,            # 77 -> 87 confirmed 2026-09-28
            "MARC -STRGR": "40", "MARC -VRMOD": 2, "MARC -VINT1": 180,
            "MARC-VINT2": 60, "MARC-MTVFP": "ST", "MARA-RAUBE": "20",
            "MBEW -BKLAS": "2000", "MBEW - PEINH": 1, "CKMLHD-MLAST": "2",
            "MBEW - VPRSV": "V", "MBEW-XLIFO": "X", "MBEW-HKMAT": "X",
            "MARC-LOSGR": 1, "MBEW-EKALR": "X", "MARC-AWSLS": "000001",
            "MARA-SERIAL": "Z001",
        },
    },
    "CA10": {
        "tab": "CA10",
        "length_unit": "CM",
        "weight_unit": "KG",
        "fixed_by_col": {107: "X"},
        "fixed": {
            "MARA-MBRSH": "M", "MARA-MTART": "Z002", "MARC-TWERK": "CA10",
            "MVKE-VKORG": "CA10", "MVKE-VTWEG": "00", "MARA-SPART": "04",
            "MARA-MEINS": "EA", "MARA-MSTAE": "03", "MARM-MEINH": "PCA",
            "MARM-UMREN": 1, "MARM-UMREZ": 1, "MARM-VOLEH": "CCM",
            "MVKE-DWERK": "CA10", "MLAN-ALAND": "CA", "TATYP": "CTXJ",
            "MLAN-TAXM1": "1", "MVKE-KTGRM": "20", "MARA-MTPOS_MARA": "NORM",
            "MVKE-MTPOS": "NORM", "MARA-TRAGR": "0001", "MARC-LADGR": "0001",
            "MARA-VABME": 1, "MARC-EKGRP": "P07",
            "MARC-WEBAZ": 4,              # 2 -> 4 confirmed 2026-09-28
            "MVKE-VRKME": "PCA",          # Sales Unit, added 2026-09-28
            "MARC-MMSTA": "03", "MARC-DISMM": "PD", "MARC-MINBE": 1,
            "MARC-DISPO": "Z02", "MARC-DISLS": "MB", "MARC-BESKZ": "F",
            "MARC-LGPRO": "C102", "MARC-LGFSB": "C102",
            "MARC -STRGR": "40", "MARC -VRMOD": 1, "MARC -VINT1": 999,
            "MARC-VINT2": 0, "MARC-MTVFP": "ST", "MARA-RAUBE": "20",
            "MARC-SERNP": "Z001", "MBEW -BKLAS": "2000", "MBEW - PEINH": 1,
            "CKMLHD-MLAST": "2", "MBEW - VPRSV": "V", "MBEW-XLIFO": "X",
            "MBEW-HKMAT": "X", "MARC-LOSGR": 1, "MBEW-EKALR": "X",
            "MARC-AWSLS": "000001",
        },
    },
}

_EU_COMMON = {
    "MARA-MBRSH": "M", "MARA-MTART": "Z002", "MVKE-VTWEG": "00",
    "MARA-SPART": "04", "MARA-MEINS": "EA", "MARA-MSTAE": "03",
    "MARM-MEINH": "ST", "MARM-UMREN": 1, "MARM-UMREZ": 1,
    "MARM-VOLEH": "CCM", "MVKE-VRKME": "ST", "MLAN-ALAND": "GB",
    "TATYP": "MWST", "MLAN-TAXM1": "1", "MVKE-KTGRM": "20",
    "MARA-MTPOS_MARA": "NORM", "MVKE-MTPOS": "NORM", "MARA-TRAGR": "0001",
    "MARC-LADGR": "0001", "MARA-VABME": 1, "MARC-EKGRP": "P07",
    "MARC-WEBAZ": 2, "MARC-MMSTA": "03", "MARC-DISPO": "Z02",
    "MARC-DISLS": "MB", "MARC-BESKZ": "F", "MARC -STRGR": "40",
    "MARC -VRMOD": 1, "MARC -VINT1": 999, "MARC-VINT2": 0,
    "MARC-MTVFP": "ST", "MARA-RAUBE": "20", "MARC-SERNP": "Z001",
    "MBEW -BKLAS": "2000", "MBEW - PEINH": 1, "CKMLHD-MLAST": "2",
    "MBEW - VPRSV": "V", "MBEW-XLIFO": "X", "MBEW-HKMAT": "X",
    "MARC-LOSGR": 1, "MBEW-EKALR": "X", "MARC-AWSLS": "000001",
}


def _eu(twerk, vkorg, dwerk, lgpro, dismm, tax_by_col, minbe=True, plifz=None,
        serial=False):
    f = dict(_EU_COMMON)
    f.update({"MARC-TWERK": twerk, "MVKE-VKORG": vkorg, "MVKE-DWERK": dwerk,
              "MARC-LGPRO": lgpro, "MARC-LGFSB": lgpro, "MARC-DISMM": dismm})
    if minbe:
        f["MARC-MINBE"] = 1
    if plifz is not None:
        f["MARC -PLIFZ"] = plifz
    if serial:
        f["MARA-SERIAL"] = "Z001"
        f.pop("MARC-SERNP", None)
        f.pop("MVKE-VRKME", None)
    return {"tab": twerk, "length_unit": "CM", "weight_unit": "KG",
            "fixed_by_col": dict(tax_by_col), "fixed": f}


PLANT_CONFIG["NL10"] = _eu("NL10", "SE10", "NL10", "NL12", "VB", EU_TAX_BY_COL)
PLANT_CONFIG["UK11"] = _eu("UK11", "SE10", "NL10", "UK11", "PD", EU_TAX_BY_COL)
PLANT_CONFIG["IT20"] = _eu("IT20", "IT10", "NL10", "IT21", "VB", EU_TAX_BY_COL)
PLANT_CONFIG["UK10"] = _eu("UK10", "UK10", "UK12", "UK10", "VB", GB_TAX_BY_COL,
                           minbe=False, plifz=77, serial=True)
PLANT_CONFIG["UK12"] = _eu("UK12", "UK10", "UK12", "UK12", "VB", GB_TAX_BY_COL,
                           minbe=False, plifz=77, serial=True)

# Order shown in the UI
PLANT_ORDER = ["US30", "CA10", "NL10", "UK11", "IT20", "UK10", "UK12"]
EU_DEFAULT = ["NL10", "UK11", "IT20"]

# Sales-Org override applied in place after generation (skill Step 6)
VKORG_OVERRIDE = {"NL10": "UK10"}
# Plants that additionally ship a Sales-Org copy: plant -> (copy sales org)
VKORG_COPY = {"US30": "US20"}

PRICE_COLUMN_FOR_PLANT = {
    "US30": "L8 US", "CA10": "L8 CA (USD)", "NL10": "L8 EUR",
    "UK11": "L8 EUR", "IT20": "L8 EUR", "UK10": "L8 EUR", "UK12": "L8 EUR",
}

# ---------------------------------------------------------------------------
# SAP field constants
# ---------------------------------------------------------------------------

SAP_MATNR = "MARA-MATNR"
SAP_BISMT = "MARA-BISMT"
SAP_DESC = "MAKT-MAKTX （EN）"      # literal space before the full-width paren
SAP_STD_PRICE = "MBEW - VERPR"
SAP_L10_COST = "MBEW-ZPLP3"
SAP_L10_DATE = "MBEW-ZPLD3"
SAP_LENGTH, SAP_WIDTH, SAP_HEIGHT = "MARM-LAENG", "MARM-BREIT", "MARM-HOEHE"
SAP_LEN_UNIT = "MARM-MEABM"
SAP_NET_WT, SAP_GROSS_WT = "MARA-NTGEW", "MARM-BRGEW"
SAP_WT_UNIT = "MARM-GEWEI"
SAP_VOLUME, SAP_VOL_UNIT = "MARM-VOLUM", "MARM-VOLEH"
SAP_HERKL = "MARC-HERKL"
SAP_PLANT = "MARC-TWERK"
SAP_VKORG = "MVKE-VKORG"

TEXT_FORMAT_SAP_FIELDS = {"MARM-EAN11", "MARC-STAWN"}

# Flat-source (Sheet1) column header -> SAP code
FLAT_FIELD_MAP = {
    "Material": SAP_MATNR,
    "Material Description": SAP_DESC,
    "Product Category": "MARA-MATKL",
    "Voltage Platform": "MARA-EXTWG",
    "EAN/UPC": "MARM-EAN11",
    "Length": SAP_LENGTH, "Width": SAP_WIDTH, "Height": SAP_HEIGHT,
    "Unit of Dimension": SAP_LEN_UNIT,
    "Volume": SAP_VOLUME, "Volume Unit": SAP_VOL_UNIT,
    "Net Weight": SAP_NET_WT, "Gross Weight": SAP_GROSS_WT,
    "Unit of Weight": SAP_WT_UNIT,
    "Sub Category": "MARA-ZMM_EXTWG_ADD",
    "Attribute1": "MARA-ZMM_EXTWG_ADD2",
    "Attribute2": "MARA-ZMM_EXTWG_ADD3",
    "Series": "MARA-ZMM_SERIES",
    "Sub-Brand": "MARA-ZMM_SUB_BRAND",
    "Warr.Periods(Y)": "MARA-ZMM_WARR_PERIODS",
    "40 HC QTY": "MARA-FERTH",
    "GW Model": "MARA-NORMT",
    "DG UN Code": "MARA-ZMM_DG_UNCODE",
    "Commodity Code": "MARC-STAWN",
    "Cntry/Reg of Origin": SAP_HERKL,
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def normalize_sap(code) -> str:
    if not code:
        return ""
    return re.sub(r"\s+", " ", str(code).strip())


def clean_str(val) -> str:
    if val is None:
        return ""
    return re.sub(r"[\xa0\s]+", " ", str(val)).strip()


def build_sap_map(ws, sap_row: int = 3) -> dict:
    mapping = {}
    row = list(ws.iter_rows(min_row=sap_row, max_row=sap_row, values_only=True))[0]
    for idx, code in enumerate(row):
        if code:
            key = normalize_sap(str(code))
            if key not in mapping:
                mapping[key] = idx
    return mapping


def _to_text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------------------
# Source readers — both return [{sap_code: value, "_row": excel_row}, ...]
# ---------------------------------------------------------------------------

def detect_format(file_bytes: bytes) -> str:
    """Return 'bulk' or 'flat'."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    first = wb[wb.sheetnames[0]]
    row3 = list(first.iter_rows(min_row=3, max_row=3, values_only=True))
    codes = [normalize_sap(c) for c in (row3[0] if row3 else []) if c]
    wb.close()
    return "bulk" if SAP_MATNR in codes else "flat"


def read_bulk_source(file_bytes: bytes):
    """Bulk 'Intake form' layout: SAP codes in row 3, data from row 10."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["Intake form"] if "Intake form" in wb.sheetnames else wb[wb.sheetnames[0]]
    sap_map = build_sap_map(ws, 3)
    col_to_sap = {v: k for k, v in sap_map.items()}
    rows = []
    for r in range(10, ws.max_row + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if not any(v is not None and str(v).strip() != "" for v in values):
            continue
        d = {}
        for idx, v in enumerate(values):
            sap = col_to_sap.get(idx)
            if sap:
                d[sap] = v
        d["_row"] = r
        d["_date"] = values[0] if values else None
        rows.append(d)
    return ws.title, rows


def read_flat_source(file_bytes: bytes):
    """Flat layout: plain headers in row 1, data from row 2."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    col_idx = {h: i for i, h in enumerate(headers) if h}
    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if not any(v is not None and str(v).strip() != "" for v in values):
            continue
        d = {}
        for header, sap in FLAT_FIELD_MAP.items():
            i = col_idx.get(header)
            d[sap] = values[i] if i is not None and i < len(values) else None
        d["_row"] = r
        d["_date"] = None
        rows.append(d)
    return ws.title, rows


# ---------------------------------------------------------------------------
# Row preparation for one target plant
# ---------------------------------------------------------------------------

def prepare_row(src: dict, plant: str) -> dict:
    """Copy a source row and apply per-row unit conversion + cleanups."""
    cfg = PLANT_CONFIG[plant]
    d = {k: v for k, v in src.items() if not k.startswith("_")}
    mat = clean_str(d.get(SAP_MATNR))
    d[SAP_BISMT] = mat or None

    converted = False

    src_len_unit = clean_str(d.get(SAP_LEN_UNIT))
    if src_len_unit and src_len_unit != cfg["length_unit"]:
        factor = INCH_TO_CM if cfg["length_unit"] == "CM" else (1 / INCH_TO_CM)
        for k in (SAP_LENGTH, SAP_WIDTH, SAP_HEIGHT):
            if d.get(k) is not None:
                try:
                    d[k] = round(float(d[k]) * factor, 2)
                except (ValueError, TypeError):
                    pass
        d[SAP_LEN_UNIT] = cfg["length_unit"]
        converted = True
    elif src_len_unit:
        d[SAP_LEN_UNIT] = cfg["length_unit"]

    src_wt_unit = clean_str(d.get(SAP_WT_UNIT)).upper()
    if src_wt_unit and src_wt_unit != cfg["weight_unit"]:
        factor = LB_TO_KG if cfg["weight_unit"] == "KG" else (1 / LB_TO_KG)
        for k in (SAP_NET_WT, SAP_GROSS_WT):
            if d.get(k) is not None:
                try:
                    d[k] = round(float(d[k]) * factor, 2)
                except (ValueError, TypeError):
                    pass
        d[SAP_WT_UNIT] = cfg["weight_unit"]
        converted = True
    elif src_wt_unit:
        d[SAP_WT_UNIT] = cfg["weight_unit"]

    # Volume is ALWAYS blanked. It is untrustworthy in both source layouts:
    # bulk files carry junk (a literal inch mark) in that column, flat files
    # carry mixed/implausible units. A stale MARM-VOLUM crashed a real SAP
    # upload on 2026-08-25. The plant's Volume Unit is a fixed template value,
    # and volume is derivable from L/W/H anyway.
    d[SAP_VOLUME] = None

    if d.get(SAP_HERKL):
        d[SAP_HERKL] = str(d[SAP_HERKL]).strip()[:2]

    for k in TEXT_FORMAT_SAP_FIELDS:
        if d.get(k) is not None:
            d[k] = _to_text(d[k])

    d["_mat"] = mat
    d["_converted"] = converted
    return d


# ---------------------------------------------------------------------------
# Output builder
# ---------------------------------------------------------------------------

def build_output(template_bytes: bytes, plant: str, src_rows: list,
                 prices: dict, out_date: date, vkorg: str | None = None) -> bytes:
    """prices: {material_code: (std_price, l10_cost)}; missing -> blank."""
    cfg = PLANT_CONFIG[plant]
    tab = cfg["tab"]
    fixed = {normalize_sap(k): v for k, v in cfg["fixed"].items()}
    fixed_by_col = cfg.get("fixed_by_col", {})
    today_str = out_date.strftime("%Y%m%d")

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    if tab not in wb.sheetnames:
        raise ValueError(f"Template has no tab named '{tab}'. "
                         f"Tabs found: {wb.sheetnames}")
    for s in list(wb.sheetnames):
        if s != tab:
            del wb[s]
    ws = wb[tab]

    sap_map = build_sap_map(ws, 3)
    col_to_sap = {v: k for k, v in sap_map.items()}
    n_cols = ws.max_column
    if ws.max_row >= 10:
        ws.delete_rows(10, ws.max_row - 9)

    text_n = {normalize_sap(f) for f in TEXT_FORMAT_SAP_FIELDS}
    text_cols = {i for i, s in col_to_sap.items() if normalize_sap(s) in text_n}

    n_std = normalize_sap(SAP_STD_PRICE)
    n_l10 = normalize_sap(SAP_L10_COST)
    n_date = normalize_sap(SAP_L10_DATE)
    n_bismt = normalize_sap(SAP_BISMT)
    n_vkorg = normalize_sap(SAP_VKORG)

    for src in src_rows:
        row = prepare_row(src, plant)
        mat = row["_mat"]
        std, l10 = prices.get(mat, (None, None))
        out = [None] * n_cols
        for ci in range(n_cols):
            sap = col_to_sap.get(ci)
            if sap is None:
                continue
            n = normalize_sap(sap)
            if n == n_vkorg and vkorg:
                out[ci] = vkorg
                continue
            if n in fixed:
                out[ci] = fixed[n]
                continue
            if n == n_date:
                out[ci] = today_str
                continue
            if n == n_bismt:
                out[ci] = mat or None
                continue
            if n == n_std:
                out[ci] = std
                continue
            if n == n_l10:
                out[ci] = l10
                continue
            if sap in row:
                out[ci] = row[sap]
        for ci, v in fixed_by_col.items():
            if ci < len(out):
                out[ci] = v
        ws.append(out)
        for cell in ws[ws.max_row]:
            cell.fill = NO_FILL
        for ci in text_cols:
            cell = ws.cell(row=ws.max_row, column=ci + 1)
            if cell.value is not None:
                cell.value = str(cell.value)
            cell.number_format = "@"

    ws.conditional_formatting._cf_rules.clear()
    for r in ws.iter_rows(min_row=10):
        for cell in r:
            cell.fill = NO_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return strip_external_links(buf.getvalue())


def make_vkorg_copy(xlsx_bytes: bytes, tab: str, new_vkorg: str) -> bytes:
    """Duplicate a finished file, changing only MVKE-VKORG on the data rows."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[tab]
    sap_map = build_sap_map(ws, 3)
    col = sap_map[normalize_sap(SAP_VKORG)] + 1
    for r in range(10, ws.max_row + 1):
        ws.cell(row=r, column=col).value = new_vkorg
    buf = io.BytesIO()
    wb.save(buf)
    return strip_external_links(buf.getvalue())


# ---------------------------------------------------------------------------
# QA: full-row contamination scan
# ---------------------------------------------------------------------------

# Fields that legitimately differ row to row — flagged but marked "expected".
EXPECTED_TO_VARY = {
    SAP_MATNR, SAP_BISMT, SAP_DESC, "MARM-EAN11", SAP_LENGTH, SAP_WIDTH,
    SAP_HEIGHT, SAP_GROSS_WT, SAP_NET_WT, "MARA-FERTH", SAP_HERKL,
    SAP_STD_PRICE, SAP_L10_COST, "MARA-NORMT", "MARC-STAWN", "MARA-MATKL",
    SAP_VOLUME,
}


# Fields SAP will try to parse as a number. A text value here is the
# CONVT_NO_NUMBER crash class (real incident 2026-08-25) — always a warning,
# even on a field that legitimately varies row to row.
NUMERIC_SAP_FIELDS = {
    SAP_LENGTH, SAP_WIDTH, SAP_HEIGHT, SAP_NET_WT, SAP_GROSS_WT, SAP_VOLUME,
    "MARA-FERTH", SAP_STD_PRICE, SAP_L10_COST,
}


def _is_numeric(v) -> bool:
    if isinstance(v, (int, float)):
        return True
    try:
        float(str(v).strip())
        return True
    except (TypeError, ValueError):
        return False


def scan_types(ws, sap_map, headers, rows):
    """Row-level sanity checks that don't depend on consensus."""
    out = []
    numeric_n = {normalize_sap(f) for f in NUMERIC_SAP_FIELDS}
    col_to_sap = {v: k for k, v in sap_map.items()}
    for c, sap in col_to_sap.items():
        n = normalize_sap(sap)
        if n not in numeric_n:
            continue
        bad = []
        for r in rows:
            v = ws.cell(row=r, column=c + 1).value
            if v is None or str(v).strip() == "":
                continue
            if not _is_numeric(v):
                bad.append((r, v))
        if bad:
            out.append({"col": c + 1, "sap": sap,
                        "header": str(headers[c] or "").replace("\n", " ")[:40],
                        "consensus": "a number",
                        "outliers": bad,
                        "kind": "non-numeric value in a numeric SAP field"})
    # country of origin must be a 2-letter code
    herkl = sap_map.get(normalize_sap(SAP_HERKL))
    if herkl is not None:
        bad = []
        for r in rows:
            v = ws.cell(row=r, column=herkl + 1).value
            if v is None or str(v).strip() == "":
                continue
            if not re.fullmatch(r"[A-Za-z]{2}", str(v).strip()):
                bad.append((r, v))
        if bad:
            out.append({"col": herkl + 1, "sap": SAP_HERKL,
                        "header": "Country of Origin", "consensus": "2 letters",
                        "outliers": bad,
                        "kind": "country of origin is not a 2-letter code"})
    return out


def scan_output(xlsx_bytes: bytes, plant: str):
    """Return (warnings, info) — warnings are columns worth a human look."""
    cfg = PLANT_CONFIG[plant]
    fixed_n = {normalize_sap(k) for k in cfg["fixed"]}
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[cfg["tab"]]
    sap_map = build_sap_map(ws, 3)
    col_to_sap = {v: k for k, v in sap_map.items()}
    headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    rows = list(range(10, ws.max_row + 1))
    warnings, info = [], []
    warnings += scan_types(ws, sap_map, headers, rows)
    if len(rows) < 2:
        return warnings, info
    for c in range(ws.max_column):
        sap = col_to_sap.get(c, "")
        if normalize_sap(sap) in fixed_n or c in cfg.get("fixed_by_col", {}):
            continue
        vals = [ws.cell(row=r, column=c + 1).value for r in rows]
        consensus, freq = Counter(vals).most_common(1)[0]
        if freq == len(vals):
            continue
        outliers = [(r, v) for r, v in zip(rows, vals) if v != consensus]
        entry = {
            "col": c + 1,
            "sap": sap or "(no SAP code)",
            "header": str(headers[c] or "").replace("\n", " ")[:40],
            "consensus": consensus,
            "outliers": outliers,
        }
        if normalize_sap(sap) in {normalize_sap(x) for x in EXPECTED_TO_VARY}:
            info.append(entry)
        else:
            warnings.append(entry)
    return warnings, info


def check_missing_prices(src_rows: list, prices: dict):
    missing = []
    for src in src_rows:
        mat = clean_str(src.get(SAP_MATNR))
        std, l10 = prices.get(mat, (None, None))
        if std is None or l10 is None:
            missing.append((mat, std, l10))
    return missing


# ---------------------------------------------------------------------------
# External-link stripping (stops Excel's "repair" prompt)
# ---------------------------------------------------------------------------

def strip_external_links(xlsx_bytes: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(xlsx_bytes), "r")
    names = zin.namelist()
    data = {n: zin.read(n) for n in names}
    zin.close()

    rels_path = "xl/_rels/workbook.xml.rels"
    wb_path = "xl/workbook.xml"
    ct_path = "[Content_Types].xml"

    rels_xml = data[rels_path].decode("utf-8")
    wb_xml = data[wb_path].decode("utf-8")
    ct_xml = data[ct_path].decode("utf-8")

    has_parts = any(n.startswith("xl/externalLinks/") for n in names)
    has_dangling = (
        "externalLink" in rels_xml
        or "<externalReferences>" in wb_xml
        or "externalLink" in ct_xml
        or re.search(r"<definedName\b[^>]*>[^<]*\[\d+\][^<]*</definedName>", wb_xml)
    )
    if not has_parts and not has_dangling:
        return xlsx_bytes

    # attribute order is not guaranteed -> lookahead matching
    rels_xml = re.sub(r'<Relationship\b(?=[^>]*Type="[^"]*externalLink")[^>]*/>', "", rels_xml)
    wb_xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", wb_xml, flags=re.DOTALL)
    wb_xml = re.sub(r"<definedName\b[^>]*>[^<]*\[\d+\][^<]*</definedName>", "", wb_xml)
    ct_xml = re.sub(
        r'<Override\b(?=[^>]*PartName="/xl/externalLinks/externalLink\d+\.xml")[^>]*/>',
        "", ct_xml)

    data[rels_path] = rels_xml.encode("utf-8")
    data[wb_path] = wb_xml.encode("utf-8")
    data[ct_path] = ct_xml.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n.startswith("xl/externalLinks/"):
                continue
            zout.writestr(n, data[n])
    return out.getvalue()


# ---------------------------------------------------------------------------
# Output naming (skill Step 8)
# ---------------------------------------------------------------------------

def output_filename(out_date: date, plant: str, vkorg: str, version: int) -> str:
    tab = PLANT_CONFIG[plant]["tab"]
    return f"{out_date.month}.{out_date.day} {tab}-{vkorg}-{version}.xlsx"


def planned_outputs(plants: list):
    """Return [(plant, sales_org, is_copy)] in delivery order."""
    out = []
    for p in plants:
        vk = VKORG_OVERRIDE.get(p, PLANT_CONFIG[p]["fixed"][SAP_VKORG])
        out.append((p, vk, False))
        if p in VKORG_COPY:
            out.append((p, VKORG_COPY[p], True))
    return out
